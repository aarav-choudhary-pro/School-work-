from pptx import Presentation
from pptx.util import Emu
import json

prs = Presentation("original/cable_tray_original.pptx")
print(f"Slide size: {prs.slide_width} x {prs.slide_height} EMU ({prs.slide_width/914400:.2f}in x {prs.slide_height/914400:.2f}in)")
print(f"Number of slides: {len(prs.slides)}")
print("="*100)

for i, slide in enumerate(prs.slides, 1):
    print(f"\n########## SLIDE {i} ##########")
    print(f"Layout: {slide.slide_layout.name}")
    for shape in slide.shapes:
        kind = shape.shape_type
        name = shape.name
        print(f"--- Shape: {name} | type={kind} | pos=({shape.left},{shape.top}) size=({shape.width},{shape.height})")
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                text = "".join(r.text for r in p.runs)
                if text.strip():
                    print(f"    TEXT: {text}")
        if shape.has_table:
            tbl = shape.table
            print(f"    TABLE {len(tbl.rows)}x{len(tbl.columns)}")
            for r_idx, row in enumerate(tbl.rows):
                cells = [c.text for c in row.cells]
                print(f"      Row{r_idx}: {cells}")
        if shape.shape_type == 13:  # picture
            print(f"    PICTURE shape")
    # notes
    if slide.has_notes_slide:
        notes = slide.notes_slide.notes_text_frame.text
        if notes.strip():
            print(f"  >> SPEAKER NOTES: {notes}")
